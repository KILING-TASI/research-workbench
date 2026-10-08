"""Export a saved human-readable financial brief to editable Word, with citations."""
import argparse,hashlib,json,re,shutil
from pathlib import Path
from urllib.parse import urlsplit

def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def blocks(text):
 result=[];lines=text.splitlines();i=0
 def cells(line):
  body=line.strip()[1:]
  if body.endswith('|'):body=body[:-1]
  return [v.strip().replace('\\|','|') for v in re.split(r'(?<!\\)\|',body)]
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  heading=re.match(r'^(#{1,3})\s+(.+)',line)
  if heading:result.append(('heading',len(heading[1]),heading[2]))
  elif line.startswith('|'):
   raw=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
   table=[cells(v) for v in raw]
   if len(table)<2 or any(not re.fullmatch(r':?-{3,}:?',v) for v in table[1]):raise ValueError('表格缺少有效分隔行')
   cols=len(table[0])
   if cols>6:raise ValueError('表格超过六列，请使用Excel底稿或拆表，避免静默压缩')
   if any(len(row)!=cols for row in table):raise ValueError('表格列数不一致，不能补零或丢弃单元格')
   result.append(('table',0,[table[0]]+table[2:]))
  else:result.append(('paragraph',0,re.sub(r'^[-*]\s+','',line)))
 return result

def export(spec):
 from docx import Document
 from docx.shared import Cm,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 source=Path(spec['markdownPath']).resolve();out=Path(spec['outDir']).resolve()
 if digest(source)!=spec['markdownSha256']:raise ValueError('报告正文哈希变化，应确认新版正文后导出')
 rows=blocks(source.read_text(encoding='utf-8-sig'))
 # Validate every target before creating an output directory.
 refs={}
 for kind,_,content in rows:
  texts=[cell for row in content for cell in row] if kind=='table' else [content]
  for text in texts:
   for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
    target=m[2];u=urlsplit(target)
    if u.scheme:
     if u.scheme not in ['http','https']:raise ValueError('不支持的引用协议')
     continue
    if not u.path:raise ValueError('空引用路径')
    p=(source.parent/u.path).resolve()
    if not p.is_relative_to(source.parent):raise ValueError('本地引用必须位于报告目录内')
    if not p.is_file():raise ValueError('引用文件缺失：'+u.path)
    refs[u.path]=p
 out.mkdir(parents=True,exist_ok=False)
 for name,path in refs.items():
  dest=out/name;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,dest)
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.3)
 for style in ['Normal','Title','Heading 1','Heading 2']:
  st=doc.styles[style];st.font.name='Microsoft YaHei';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
 doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
 # Word's built-in Title style may carry a theme border; keep titles unruled.
 for style in ['Title','Heading 1','Heading 2']:
  ppr=doc.styles[style].element.find(qn('w:pPr'))
  if ppr is not None:
   for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
  fonts=doc.styles[style].element.rPr.rFonts
  for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   fonts.attrib.pop(qn('w:'+attribute),None)
  fonts.set(qn('w:ascii'),'Microsoft YaHei');fonts.set(qn('w:hAnsi'),'Microsoft YaHei')
 def inline(p,text):
  start=0
  for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
   p.add_run(text[start:m.start()].replace('**',''))
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245A81');prop.append(color);run.append(prop);t=OxmlElement('w:t');t.text=m[1];run.append(t);h.append(run);p._p.append(h);start=m.end()
  p.add_run(text[start:].replace('**',''))
 for kind,level,text in rows:
  if kind=='table':
   table=doc.add_table(rows=len(text),cols=len(text[0]));table.style='Table Grid'
   for i,row in enumerate(text):
    tr=table.rows[i]._tr;prop=tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
    if i==0:prop.append(OxmlElement('w:tblHeader'))
    for j,value in enumerate(row):
     para=table.cell(i,j).paragraphs[0];inline(para,value)
     for run in para.runs:run.font.size=Pt(9);run.bold=i==0
   doc.add_paragraph('')
  else:
   p=doc.add_paragraph(style=('Title' if level==1 else 'Heading '+str(level-1)) if kind=='heading' else 'Normal');inline(p,text)
   if kind=='heading':p.paragraph_format.keep_with_next=True
 doc.add_paragraph('仅供公开资料研究与复核，不构成投资建议。缺失资料及口径限制以正文说明为准。')
 path=out/'公司与行业点评.docx';doc.save(path)
 result={'outputFile':path.name,'format':'docx','markdownSha256':digest(source),'outputSha256':digest(path),'paragraphs':sum(kind!='table' for kind,_,_ in rows),'tables':sum(kind=='table' for kind,_,_ in rows),'references':[{'path':name,'sha256':digest(p)} for name,p in sorted(refs.items())],'limitations':['导出只保留已存正文和引用，不新增或独立核验研究结论','源正文哈希确认不代表正文与数据快照语义一致','排版需渲染复核，保存成功不等于视觉验收通过','支持六列以内规则Markdown表格，不处理合并单元格或嵌套表；复杂财务模型使用Excel底稿']}
 (out/'export-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');return result
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('input');args=p.parse_args();print(json.dumps(export(json.loads(Path(args.input).read_text(encoding='utf-8-sig'))),ensure_ascii=False))
